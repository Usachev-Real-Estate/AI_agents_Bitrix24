"""Import call-center owner leads from Excel into Bitrix24 contact + seller deal."""

from __future__ import annotations

import json
import logging
import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from config import Settings
from notify import _bx_call_sync

logger = logging.getLogger(__name__)

MSK = ZoneInfo("Europe/Moscow")
DEFAULT_EXCEL = Path("data/лиды_успешно добавить в битрикс собы от кц.xlsx")
DEFAULT_SHEET = "Успешно"
IMPORT_LOG = Path("data/kc_import_results.jsonl")

CONTACT_TYPE_OWNER = "UC_2G0TD3"
SOURCE_KC = "24"
DEAL_CATEGORY_ID = 0
DEAL_STAGE_NEW = "NEW"
PROPERTY_TYPE_APARTMENT = "356"
UF_OBJECT_COST = "UF_CRM_1774521607469"
UF_PROPERTY_TYPE = "UF_CRM_1747291787883"
UF_ZK_NAME = "UF_CRM_1774364892961"

BROKER_USER_IDS: dict[str, int] = {
    "Пьянкова Ульяна": 66,
    "Ирена Азизова": 190,
    "Ветров Владислав": 82,
    "Абзалилов Марат": 96,
    "Галина Шарипова": 78,
}

EXCEL_HEADERS = (
    "ID",
    "Hash",
    "Имя",
    "Телефон",
    "ЗК",
    "Комментарий",
    "Создан",
    "Обновлён",
    "Bitrix24 ID",
    "Синхр. Bitrix24",
    "Оператор",
    "Статус",
    "ответственный сотрудник",
)

CREATED_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})\s+(\d{2}):(\d{2})$")


@dataclass(frozen=True)
class KcRow:
    """One importable row from the KC Excel sheet."""

    row_num: int
    external_id: str
    row_hash: str
    name: str
    phone: str
    zk: str
    comment: str
    created_raw: str
    responsible: str
    existing_bitrix_id: str = ""


@dataclass
class ImportResult:
    """Outcome of importing one Excel row."""

    row_num: int
    external_id: str
    row_hash: str
    status: str
    contact_id: int | None = None
    deal_id: int | None = None
    error: str = ""


def normalize_phone(raw: Any) -> str:
    """Normalize phone to +7XXXXXXXXXX."""
    digits = re.sub(r"\D", "", str(raw or ""))
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    if len(digits) == 10:
        digits = "7" + digits
    if len(digits) != 11 or not digits.startswith("7"):
        raise ValueError(f"invalid phone: {raw!r}")
    return f"+{digits}"


def parse_created_at(raw: str) -> str:
    """Parse Excel datetime DD.MM.YYYY HH:MM to ISO in MSK."""
    text = (raw or "").strip()
    match = CREATED_RE.match(text)
    if not match:
        raise ValueError(f"invalid created date: {raw!r}")
    day, month, year, hour, minute = map(int, match.groups())
    dt = datetime(year, month, day, hour, minute, tzinfo=MSK)
    return dt.isoformat()


def resolve_broker_id(responsible: str) -> int:
    """Map Excel responsible name to Bitrix user ID."""
    name = (responsible or "").strip()
    if name not in BROKER_USER_IDS:
        raise ValueError(f"unknown responsible broker: {name!r}")
    return BROKER_USER_IDS[name]


def build_contact_fields(row: KcRow) -> dict[str, Any]:
    """Build crm.contact.add fields for one row."""
    assigned_id = resolve_broker_id(row.responsible)
    comment_lines: list[str] = []
    if row.created_raw.strip():
        comment_lines.append(f"Дата создания в КЦ: {row.created_raw.strip()}")
    if row.zk.strip():
        comment_lines.append(row.zk.strip())
    if row.comment.strip():
        comment_lines.append(row.comment.strip())
    return {
        "NAME": row.name.strip(),
        "LAST_NAME": row.external_id.strip(),
        "PHONE": [{"VALUE": normalize_phone(row.phone), "VALUE_TYPE": "WORK"}],
        "ASSIGNED_BY_ID": assigned_id,
        "TYPE_ID": CONTACT_TYPE_OWNER,
        "SOURCE_ID": SOURCE_KC,
        "COMMENTS": "\n".join(comment_lines),
    }


def build_deal_fields(row: KcRow, contact_id: int) -> dict[str, Any]:
    """Build crm.deal.add fields for one row."""
    assigned_id = resolve_broker_id(row.responsible)
    zk = row.zk.strip()
    ext_id = row.external_id.strip()
    return {
        "TITLE": f"{zk} {ext_id}".strip(),
        "CATEGORY_ID": DEAL_CATEGORY_ID,
        "STAGE_ID": DEAL_STAGE_NEW,
        "SOURCE_ID": SOURCE_KC,
        "ASSIGNED_BY_ID": assigned_id,
        "CONTACT_ID": contact_id,
        "OPPORTUNITY": 0,
        "IS_MANUAL_OPPORTUNITY": "Y",
        UF_OBJECT_COST: "0|RUB",
        UF_PROPERTY_TYPE: PROPERTY_TYPE_APARTMENT,
        UF_ZK_NAME: [zk],
    }


def load_rows_from_excel(path: Path, sheet: str = DEFAULT_SHEET) -> list[KcRow]:
    """Read import rows from Excel; skip empty name rows."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required: pip install openpyxl") from exc

    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found in {path}")
    ws = wb[sheet]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if headers[: len(EXCEL_HEADERS)] != list(EXCEL_HEADERS):
        logger.warning("Unexpected headers: %s", headers)

    rows: list[KcRow] = []
    for idx, excel_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(excel_row) + [None] * max(0, len(EXCEL_HEADERS) - len(excel_row))
        name = str(values[2] or "").strip()
        if not name:
            continue
        rows.append(
            KcRow(
                row_num=idx,
                external_id=str(values[0] or "").strip(),
                row_hash=str(values[1] or "").strip(),
                name=name,
                phone=str(values[3] or "").strip(),
                zk=str(values[4] or "").strip(),
                comment=str(values[5] or "").strip(),
                created_raw=str(values[6] or "").strip(),
                responsible=str(values[12] or "").strip(),
                existing_bitrix_id=str(values[8] or "").strip(),
            )
        )
    return rows


def load_import_log(path: Path = IMPORT_LOG) -> dict[str, ImportResult]:
    """Load prior import results keyed by external_id."""
    if not path.exists():
        return {}
    results: dict[str, ImportResult] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        payload = json.loads(line)
        ext_id = str(payload.get("external_id") or "")
        if not ext_id:
            continue
        results[ext_id] = ImportResult(
            row_num=int(payload.get("row_num") or 0),
            external_id=ext_id,
            row_hash=str(payload.get("row_hash") or ""),
            status=str(payload.get("status") or ""),
            contact_id=int(payload["contact_id"]) if payload.get("contact_id") else None,
            deal_id=int(payload["deal_id"]) if payload.get("deal_id") else None,
            error=str(payload.get("error") or ""),
        )
    return results


def append_import_log(result: ImportResult, path: Path = IMPORT_LOG) -> None:
    """Append one import result as JSONL."""
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "row_num": result.row_num,
        "external_id": result.external_id,
        "row_hash": result.row_hash,
        "status": result.status,
        "contact_id": result.contact_id,
        "deal_id": result.deal_id,
        "error": result.error,
    }
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(payload, ensure_ascii=False) + "\n")


def add_timeline_comment(
    entity_type: str,
    entity_id: int,
    comment: str,
    *,
    dry_run: bool,
) -> None:
    """Add timeline comment to contact or deal."""
    text = (comment or "").strip()
    if not text:
        return
    if dry_run:
        logger.info(
            "DRY_RUN: would add timeline comment to %s #%s",
            entity_type,
            entity_id,
        )
        return
    _bx_call_sync(
        "crm.timeline.comment.add",
        {
            "fields": {
                "ENTITY_TYPE": entity_type,
                "ENTITY_ID": entity_id,
                "COMMENT": text,
            }
        },
    )


def import_row(row: KcRow, *, dry_run: bool, pause_sec: float = 0.35) -> ImportResult:
    """Create contact, deal, and timeline comments for one Excel row."""
    try:
        contact_fields = build_contact_fields(row)
        deal_preview = build_deal_fields(row, contact_id=0)
    except ValueError as exc:
        return ImportResult(
            row_num=row.row_num,
            external_id=row.external_id,
            row_hash=row.row_hash,
            status="error",
            error=str(exc),
        )

    if dry_run:
        logger.info(
            "DRY_RUN row %s id=%s: contact=%s deal=%s",
            row.row_num,
            row.external_id,
            contact_fields,
            {**deal_preview, "CONTACT_ID": "<new>"},
        )
        return ImportResult(
            row_num=row.row_num,
            external_id=row.external_id,
            row_hash=row.row_hash,
            status="dry_run_skipped",
        )

    try:
        contact_id = int(
            _bx_call_sync(
                "crm.contact.add",
                {"fields": contact_fields},
            )
        )
        add_timeline_comment("contact", contact_id, row.comment, dry_run=False)
        time.sleep(pause_sec)

        deal_fields = build_deal_fields(row, contact_id)
        deal_id = int(_bx_call_sync("crm.deal.add", {"fields": deal_fields}))
        add_timeline_comment("deal", deal_id, row.comment, dry_run=False)
        time.sleep(pause_sec)

        logger.info(
            "Imported row %s id=%s -> contact #%s deal #%s",
            row.row_num,
            row.external_id,
            contact_id,
            deal_id,
        )
        return ImportResult(
            row_num=row.row_num,
            external_id=row.external_id,
            row_hash=row.row_hash,
            status="ok",
            contact_id=contact_id,
            deal_id=deal_id,
        )
    except Exception as exc:
        logger.exception("Failed row %s id=%s", row.row_num, row.external_id)
        return ImportResult(
            row_num=row.row_num,
            external_id=row.external_id,
            row_hash=row.row_hash,
            status="error",
            error=str(exc),
        )


def update_excel_bitrix_ids(
    path: Path,
    results: list[ImportResult],
    sheet: str = DEFAULT_SHEET,
) -> int:
    """Write contact/deal IDs back to Bitrix24 ID column for successful rows."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required") from exc

    ok_by_row = {
        r.row_num: f"{r.contact_id}/{r.deal_id}"
        for r in results
        if r.status == "ok" and r.contact_id and r.deal_id
    }
    if not ok_by_row:
        return 0

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    updated = 0
    for row_num, value in ok_by_row.items():
        ws.cell(row=row_num, column=9, value=value)
        ws.cell(row=row_num, column=10, value=datetime.now(MSK).strftime("%d.%m.%Y %H:%M"))
        updated += 1
    wb.save(path)
    return updated


def run_import(
    settings: Settings,
    *,
    excel_path: Path = DEFAULT_EXCEL,
    sheet: str = DEFAULT_SHEET,
    limit: int | None = None,
    dry_run: bool | None = None,
    skip_excel_update: bool = False,
    pause_sec: float = 0.35,
) -> dict[str, Any]:
    """Import all eligible rows from Excel."""
    effective_dry_run = settings.dry_run if dry_run is None else dry_run
    rows = load_rows_from_excel(excel_path, sheet=sheet)
    prior = load_import_log()

    to_process: list[KcRow] = []
    skipped = 0
    for row in rows:
        if row.existing_bitrix_id:
            skipped += 1
            continue
        if row.external_id in prior and prior[row.external_id].status == "ok":
            skipped += 1
            continue
        to_process.append(row)

    if limit is not None:
        to_process = to_process[:limit]

    results: list[ImportResult] = []
    for row in to_process:
        result = import_row(row, dry_run=effective_dry_run, pause_sec=pause_sec)
        results.append(result)
        if not effective_dry_run and result.status == "ok":
            append_import_log(result)

    ok = sum(1 for r in results if r.status == "ok")
    dry_skipped = sum(1 for r in results if r.status == "dry_run_skipped")
    errors = [r for r in results if r.status == "error"]

    excel_updated = 0
    if not effective_dry_run and not skip_excel_update and ok:
        excel_updated = update_excel_bitrix_ids(excel_path, results, sheet=sheet)

    summary = {
        "excel_path": str(excel_path),
        "dry_run": effective_dry_run,
        "total_rows_in_file": len(rows),
        "skipped_already_imported": skipped,
        "processed": len(results),
        "ok": ok,
        "dry_run_skipped": dry_skipped,
        "errors": len(errors),
        "excel_updated": excel_updated,
        "error_details": [
            {"row_num": r.row_num, "external_id": r.external_id, "error": r.error}
            for r in errors
        ],
    }
    logger.info("KC import finished: %s", summary)
    return summary
